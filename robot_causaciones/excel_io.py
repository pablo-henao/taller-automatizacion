"""Lectura de la cola PorRadicar, limpieza/validación de datos y escritura de Resultados."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNAS_ENTRADA = ["IdLista", "Factura", "NIT", "Proveedor", "Valor", "CentroCosto", "FechaFactura"]
COLUMNAS_SALIDA = ["IdLista", "Factura", "NIT", "Estado", "Causacion", "Detalle"]


@dataclass
class Fila:
    id_lista: int
    factura: str
    nit: str
    proveedor: str
    valor: object
    centro: str
    fecha: object


@dataclass
class Resultado:
    id_lista: int
    factura: str
    nit: str
    estado: str
    causacion: str = ""
    detalle: str = ""


def _texto(valor: object) -> str:
    """Texto limpio de una celda; un número entero llega sin '.0'."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def limpiar_nit(valor: object) -> str:
    """900123456, 900123456.0, '900,123,456' y '900.123.456' -> '900123456'."""
    return re.sub(r"[.,\s]", "", _texto(valor)) if not isinstance(valor, float) else _texto(valor)


def parsear_id(valor: object) -> int | None:
    texto = _texto(valor)
    return int(texto) if re.fullmatch(r"\d+", texto) else None


def parsear_valor(valor: object) -> int | None:
    """Devuelve el valor como entero > 0, o None si no es un entero positivo."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        numero = int(valor) if float(valor).is_integer() else None
    elif isinstance(valor, str):
        texto = valor.strip()
        numero = int(texto.split(".")[0]) if re.fullmatch(r"\d+(\.0+)?", texto) else None
    else:
        numero = None
    return numero if numero is not None and numero > 0 else None


def parsear_fecha(valor: object) -> date | None:
    """Acepta fechas de Excel o texto yyyy-MM-dd; None si no es una fecha válida."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", valor.strip()):
        try:
            return datetime.strptime(valor.strip(), "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def leer_cola(ruta: Path) -> list[Fila]:
    libro = load_workbook(ruta, data_only=True)
    hoja = libro.worksheets[0]
    filas = list(hoja.iter_rows(values_only=True))

    # La fila de encabezados es la primera que contiene IdLista (normalmente la 1).
    indice_encabezado = next(
        (i for i, f in enumerate(filas[:10]) if any(_texto(c).lower() == "idlista" for c in f)), None
    )
    if indice_encabezado is None:
        raise ValueError(f"No se encontró la columna IdLista en la primera hoja de {ruta}")
    posicion = {_texto(c).lower(): j for j, c in enumerate(filas[indice_encabezado]) if _texto(c)}
    faltan = [c for c in COLUMNAS_ENTRADA if c.lower() not in posicion]
    if faltan:
        raise ValueError(f"Faltan columnas en {ruta}: {', '.join(faltan)}")

    def celda(fila: tuple, nombre: str) -> object:
        j = posicion[nombre.lower()]
        return fila[j] if j < len(fila) else None

    cola: list[Fila] = []
    for fila in filas[indice_encabezado + 1 :]:
        id_lista = parsear_id(celda(fila, "IdLista"))
        if id_lista is None:  # filas sin IdLista se ignoran
            continue
        cola.append(
            Fila(
                id_lista=id_lista,
                factura=_texto(celda(fila, "Factura")),
                nit=limpiar_nit(celda(fila, "NIT")),
                proveedor=_texto(celda(fila, "Proveedor")),
                valor=celda(fila, "Valor"),
                centro=_texto(celda(fila, "CentroCosto")),
                fecha=celda(fila, "FechaFactura"),
            )
        )
    return cola


def escribir_resultados(ruta: Path, resultados: list[Resultado]) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resultados"
    hoja.append(COLUMNAS_SALIDA)
    for r in resultados:
        nit = int(r.nit) if r.nit.isdigit() else r.nit
        hoja.append([r.id_lista, r.factura, nit, r.estado, r.causacion or None, r.detalle or None])

    ultima_fila = max(len(resultados) + 1, 2)  # una tabla necesita al menos una fila de datos
    tabla = Table(displayName="Resultados", ref=f"A1:{get_column_letter(len(COLUMNAS_SALIDA))}{ultima_fila}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    hoja.add_table(tabla)
    for j, ancho in enumerate([10, 14, 14, 16, 18, 60], start=1):
        hoja.column_dimensions[get_column_letter(j)].width = ancho
    libro.save(ruta)
